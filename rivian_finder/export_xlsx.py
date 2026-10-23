"""Excel workbook export: formatted tables, frozen headers, filters, currency formats, conditional formatting."""
from __future__ import annotations

import io
import sqlite3
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import finance, hiding, linkcheck, queries, warranty
from .config import TONNEAU_STATUSES

MONEY = '"$"#,##0'
MONEY_CENTS = '"$"#,##0.00'
MILES = "#,##0"
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
GREY = PatternFill("solid", fgColor="EDEDED")


def _write_table(ws, df: pd.DataFrame, name: str, money: tuple = (), cents: tuple = (), miles: tuple = (),
                 links: tuple = (), widths: dict | None = None, style: str = "TableStyleMedium2") -> None:
    """Write a DataFrame as a native Excel table (filter buttons included) with a frozen header row."""
    cols = list(df.columns) if not df.empty else list(df.columns) or ["(no rows)"]
    ws.append(cols)
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) or v is pd.NA else
                   (bool(v) if isinstance(v, (bool,)) else v) for v in row])
    n_rows = max(len(df), 1)
    if df.empty:
        ws.append([""] * len(cols))
    ref = f"A1:{get_column_letter(len(cols))}{n_rows + 1}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "B2" if len(cols) > 3 else "A2"
    for idx, col in enumerate(cols, start=1):
        letter = get_column_letter(idx)
        fmt = MONEY if col in money else MONEY_CENTS if col in cents else MILES if col in miles else None
        for cell in ws[letter][1:]:
            if fmt:
                cell.number_format = fmt
            if col in links and cell.value:
                cell.hyperlink = str(cell.value).split(" ")[0]
                cell.font = Font(color="0563C1", underline="single")
        width = (widths or {}).get(col)
        if width is None:
            sample = [len(str(col))] + [len(str(c.value)) for c in ws[letter][1:200] if c.value is not None]
            width = min(max(sample) + 2, 48)
        ws.column_dimensions[letter].width = width
    for cell in ws[1]:
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30


def _col(ws, header: str) -> str | None:
    for cell in ws[1]:
        if cell.value == header:
            return cell.column_letter
    return None


def _now_local() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def build_workbook(conn: sqlite3.Connection, settings: dict) -> Workbook:
    inv_all = queries.inventory_frame(conn, settings)
    filters = settings["filters"]
    if settings["export"].get("include_filtered_out"):
        inv = inv_all[~inv_all["hidden"]] if not inv_all.empty else inv_all
    else:
        inv = queries.apply_filters(inv_all, filters)
    wb = Workbook()

    # 1. Current Inventory
    ws = wb.active
    ws.title = "Current Inventory"
    cur_cols = {
        "fit_score": "Fit Score", "favorite": "Favorite", "year": "Year", "make": "Brand", "model": "Model",
        "trim": "Trim", "generation": "Generation", "price": "Price", "your_price": "Your Price (after discounts)",
        "adjusted_price": "Adjusted Price (your values)", "adjust_note": "Adjustments", "miles": "Mileage",
        "price_change": "Price Change Since First Seen",
        "market_price": "Market Estimate", "deal_delta": "vs Market ($)", "deal_label": "Deal Rating",
        "comps_text": "Comparable Trucks", "days_listed": "Days Listed", "price_drops": "Price Drops", "dropped_total": "Total Dropped",
        "leverage": "Negotiation Leverage (heuristic)", "battery": "Battery", "range_text": "EPA Range (mi)",
        "battery_src": "Battery Source", "motor": "Motor", "motor_src": "Motor Source", "color": "Color",
        "tow_capacity_lbs": "Tow Capacity (lb)", "tow_src": "Tow Source", "tow_conf": "Tow Confidence",
        "warranty_ends": "Warranty Ends", "warranty": "Warranty Remaining",
        "subframe": "Front Subframe (est.)",
        "seller": "Seller", "location": "Location", "distance_mi": "Distance from Home (mi)", "source": "Source",
        "sources": "All Sources", "source_count": "Listing Count", "vin": "VIN",
        "sold_by": "Sold By", "seller_site": "Seller Type / Site", "seller_url": "Seller Listing URL",
        "link_label": "Link Check", "found_on": "Found On (aggregators)",
        "found_on_url": "Aggregator URL", "find_seller_url": "Find Seller (web search)", "url": "Listing URL",
        "history_report_url": "Carfax Report URL", "photo_url": "Photo URL",
        "transport_cost": "Transport Estimate", "transport_basis": "Transport Basis", "flags": "History Flags",
        "flip_text": "History Flag Details", "seller_phone": "Seller Phone",
        "first_seen": "First Seen (UTC)", "last_seen": "Captured / Last Seen (UTC)", "status": "Status",
        "est_monthly": "Est. Monthly (primary scenario)", "cash_to_target": "Cash Needed for Target Payment",
        "total_cost": "Approx. Total Acquisition Cost", "conflicts": "Source Conflicts", "notes": "Notes",
    }
    if not inv.empty:
        inv = inv.assign(link_label=inv["link_result"].map(linkcheck.label), warranty_ends=[warranty.ends_text(e, m) for e, m in zip(inv["warranty_expires"],
                                                                                  inv["warranty_miles_limit"])])
    cur = inv[list(cur_cols)].rename(columns=cur_cols) if not inv.empty else pd.DataFrame(columns=list(cur_cols.values()))
    _write_table(ws, cur, "CurrentInventory",
                 money=("Price", "Your Price (after discounts)", "Adjusted Price (your values)", "Price Change Since First Seen", "Approx. Total Acquisition Cost",
                        "Cash Needed for Target Payment", "Total Dropped", "Market Estimate", "vs Market ($)",
                        "Transport Estimate"),
                 cents=("Est. Monthly (primary scenario)",), miles=("Mileage", "Distance from Home (mi)"),
                 links=("Listing URL", "Seller Listing URL", "Aggregator URL", "Find Seller (web search)",
                        "Carfax Report URL", "Photo URL"),
                 widths={"Listing URL": 40, "Seller Listing URL": 40, "Aggregator URL": 30,
                         "Find Seller (web search)": 22, "Notes": 30, "Source Conflicts": 20})
    last = len(cur) + 1
    if not cur.empty:
        c = _col(ws, "vs Market ($)")
        ws.conditional_formatting.add(f"{c}2:{c}{last}", CellIsRule(operator="lessThan", formula=["0"], fill=GREEN))
        ws.conditional_formatting.add(f"{c}2:{c}{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=RED))
        c = _col(ws, "Price Change Since First Seen")
        ws.conditional_formatting.add(f"{c}2:{c}{last}", CellIsRule(operator="lessThan", formula=["0"], fill=GREEN))
        ws.conditional_formatting.add(f"{c}2:{c}{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=RED))
        c = _col(ws, "Est. Monthly (primary scenario)")
        target = settings["finance"]["target_monthly"]
        ws.conditional_formatting.add(f"{c}2:{c}{last}", CellIsRule(operator="lessThanOrEqual", formula=[str(target)], fill=GREEN))

    # 2. New Listings
    days = int(settings["export"].get("new_listing_window_days", 7))
    since = (datetime.now(timezone.utc) - timedelta(days=days)).strftime("%Y-%m-%dT%H:%M:%SZ")
    ev = queries.events_frame(conn, since=since, types=["new_listing", "new_source_for_vin", "returned"])
    ws = wb.create_sheet("New Listings")
    ev_cols = {"created_at": "Detected (UTC)", "event": "Event", "year": "Year", "model": "Model",
               "current_price": "Current Price", "current_miles": "Mileage", "seller_name": "Seller",
               "state": "State", "source": "Source", "vin": "VIN", "url": "Listing URL", "is_demo": "Demo Data"}
    _write_table(ws, _subset(ev, ev_cols), "NewListings", money=("Current Price",), miles=("Mileage",),
                 links=("Listing URL",), widths={"Listing URL": 40})

    # 3. Price Changes
    pc = queries.events_frame(conn, types=["price_drop", "price_increase"])
    if not pc.empty:
        pc["old_price"] = pd.to_numeric(pc["old_value"], errors="coerce")
        pc["new_price"] = pd.to_numeric(pc["new_value"], errors="coerce")
        pc["delta"] = pc["new_price"] - pc["old_price"]
        pc["pct"] = (pc["delta"] / pc["old_price"]).round(4)
    ws = wb.create_sheet("Price Changes")
    pc_cols = {"created_at": "Detected (UTC)", "event": "Change", "year": "Year", "model": "Model", "old_price": "Old Price",
               "new_price": "New Price", "delta": "Change ($)", "pct": "Change (%)", "current_miles": "Mileage",
               "seller_name": "Seller", "state": "State", "source": "Source", "vin": "VIN", "url": "Listing URL",
               "is_demo": "Demo Data"}
    pcd = _subset(pc, pc_cols)
    _write_table(ws, pcd, "PriceChanges", money=("Old Price", "New Price", "Change ($)"), miles=("Mileage",),
                 links=("Listing URL",), widths={"Listing URL": 40})
    if not pcd.empty:
        c, p = _col(ws, "Change ($)"), _col(ws, "Change (%)")
        ws.conditional_formatting.add(f"{c}2:{c}{len(pcd) + 1}", CellIsRule(operator="lessThan", formula=["0"], fill=GREEN))
        ws.conditional_formatting.add(f"{c}2:{c}{len(pcd) + 1}", CellIsRule(operator="greaterThan", formula=["0"], fill=RED))
        for cell in ws[p][1:]:
            cell.number_format = "0.0%"

    # 4. Missing / Removed
    mr = pd.read_sql_query(
        "SELECT l.status AS Status, l.status_reason AS Reason, l.missing_since AS \"Missing Since (UTC)\","
        " l.last_seen_at AS \"Last Seen (UTC)\", v.year AS Year, v.make AS Brand, v.model AS Model, l.current_price AS \"Last Price\","
        " l.current_miles AS Mileage, l.seller_name AS Seller, l.state AS State, l.source_key AS Source, v.vin AS VIN,"
        " l.url AS \"Listing URL\", l.is_demo AS \"Demo Data\" FROM listings l JOIN vehicles v ON v.id = l.vehicle_id"
        " WHERE l.status IN ('missing','removed') ORDER BY l.missing_since DESC", conn)
    if not mr.empty:
        mr["Source"] = mr["Source"].map(queries.source_label)
        mr["Demo Data"] = mr["Demo Data"].astype(bool)
    ws = wb.create_sheet("Missing - Removed")
    _write_table(ws, mr, "MissingRemoved", money=("Last Price",), miles=("Mileage",), links=("Listing URL",),
                 widths={"Listing URL": 40, "Reason": 45})
    if not mr.empty:
        c = _col(ws, "Status")
        ws.conditional_formatting.add(f"{c}2:{c}{len(mr) + 1}", CellIsRule(operator="equal", formula=['"missing"'], fill=YELLOW))
        ws.conditional_formatting.add(f"{c}2:{c}{len(mr) + 1}", CellIsRule(operator="equal", formula=['"removed"'], fill=GREY))

    # 5. Listing History
    hist = pd.read_sql_query(
        "SELECT o.observed_at AS \"Captured (UTC)\", v.year AS Year, v.make AS Brand, v.model AS Model, o.price AS Price, o.miles AS Mileage,"
        " o.available AS Available, o.import_method AS Method, l.source_key AS Source, o.seller_name AS Seller,"
        " v.vin AS VIN, l.url AS \"Listing URL\", o.run_id AS Run, l.is_demo AS \"Demo Data\""
        " FROM observations o JOIN listings l ON l.id = o.listing_id JOIN vehicles v ON v.id = l.vehicle_id"
        " ORDER BY v.id, l.id, o.observed_at", conn)
    if not hist.empty:
        hist["Source"] = hist["Source"].map(queries.source_label)
        hist["Available"] = hist["Available"].astype(bool)
        hist["Demo Data"] = hist["Demo Data"].astype(bool)
    ws = wb.create_sheet("Listing History")
    _write_table(ws, hist, "ListingHistory", money=("Price",), miles=("Mileage",), links=("Listing URL",),
                 widths={"Listing URL": 40})

    # 6. Source Status
    ss = queries.source_status(conn)
    ss = ss.rename(columns={"key": "Key", "name": "Source", "kind": "Type", "enabled": "Enabled",
                            "last_attempt_at": "Last Attempt (UTC)", "last_success_at": "Last Success (UTC)",
                            "last_status": "Last Status", "last_error": "Last Error", "active_listings": "Active Listings",
                            "missing_listings": "Missing Listings", "last_notes": "Notes", "last_complete": "Last Scan Complete"})
    ss["Enabled"] = ss["Enabled"].astype(bool)
    ss["Last Scan Complete"] = ss["Last Scan Complete"].map(lambda v: None if pd.isna(v) else bool(v))
    ws = wb.create_sheet("Source Status")
    _write_table(ws, ss, "SourceStatus", widths={"Notes": 60, "Last Error": 40})
    c = _col(ws, "Last Status")
    n = len(ss) + 1
    ws.conditional_formatting.add(f"{c}2:{c}{n}", CellIsRule(operator="equal", formula=['"success"'], fill=GREEN))
    ws.conditional_formatting.add(f"{c}2:{c}{n}", CellIsRule(operator="equal", formula=['"partial"'], fill=YELLOW))
    ws.conditional_formatting.add(f"{c}2:{c}{n}", CellIsRule(operator="equal", formula=['"failed"'], fill=RED))

    # 7. Financial Comparisons
    fa = finance.FinanceAssumptions.from_settings(settings["finance"])
    rows = []
    for r in inv.itertuples(index=False):
        if not r.price:
            continue
        for sc in finance.all_scenarios(float(r.price), fa, float(r.transport_cost or 0)):
            row = {"Year": r.year, "Brand": r.make, "Model": r.model, "Trim": r.trim, "Battery": r.battery, "Motor": r.motor,
                   "Powered Tonneau": TONNEAU_STATUSES.get(r.tonneau, r.tonneau), "Price": r.price,
                   "Scenario": sc.label, "Credit (trade/cash)": sc.credit_value, "Loan Payoff": sc.payoff,
                   "Equity": sc.equity, "Taxable Amount": sc.taxable_amount, "Sales Tax": sc.sales_tax,
                   "DMV + Doc Fees": sc.fees, "Transport": sc.transport,
                   "Total Acquisition Cost": sc.total_acquisition_cost, "Amount Financed": sc.amount_financed}
            for k, pay in sc.payments.items():
                row[f"Payment {k.replace('@', ' mo @ ')}%"] = pay
            for k, cash in sc.cash_to_target.items():
                row[f"Cash for ${fa.target_monthly:,.0f}/mo {k.replace('@', ' mo @ ')}%"] = cash
            row.update({"Seller": r.seller, "Source": r.source, "VIN": r.vin, "Listing URL": r.url,
                        "Demo Data": r.is_demo})
            rows.append(row)
    fin_df = pd.DataFrame(rows)
    ws = wb.create_sheet("Financial Comparisons")
    money_cols = tuple(c for c in fin_df.columns if c in (
        "Price", "Credit (trade/cash)", "Loan Payoff", "Equity", "Taxable Amount", "DMV + Doc Fees", "Transport",
        "Total Acquisition Cost", "Amount Financed") or c.startswith("Cash for"))
    cents_cols = tuple(c for c in fin_df.columns if c.startswith("Payment") or c == "Sales Tax")
    _write_table(ws, fin_df, "FinancialComparisons", money=money_cols, cents=cents_cols, links=("Listing URL",),
                 widths={"Listing URL": 40, "Scenario": 34})
    if not fin_df.empty:
        for col in fin_df.columns:
            if col.startswith("Payment"):
                c = _col(ws, col)
                ws.conditional_formatting.add(f"{c}2:{c}{len(fin_df) + 1}", CellIsRule(
                    operator="lessThanOrEqual", formula=[str(fa.target_monthly)], fill=GREEN))
                ws.conditional_formatting.add(f"{c}2:{c}{len(fin_df) + 1}", CellIsRule(
                    operator="greaterThan", formula=[str(fa.target_monthly)], fill=RED))

    # Not interested (so a shared workbook explains why trucks are missing from the list)
    hidden = inv_all[inv_all["hidden"]] if not inv_all.empty else inv_all
    hid_cols = {"hidden_at": "Hidden (UTC)", "hidden_reason": "Reason", "hidden_note": "Note", "year": "Year", "make": "Brand", "model": "Model",
                "trim": "Trim", "price": "Price", "miles": "Mileage", "title_status": "Title",
                "accident_history": "Accidents", "tonneau_label": "Powered Tonneau", "seller": "Seller",
                "location": "Location", "vin": "VIN", "seller_url": "Seller Listing URL", "found_on_url": "Aggregator URL",
                "is_demo": "Demo Data"}
    hd = hidden[list(hid_cols)].rename(columns=hid_cols) if not hidden.empty else pd.DataFrame(columns=list(hid_cols.values()))
    if not hd.empty:
        hd["Reason"] = hd["Reason"].map(hiding.reason_label)
    ws = wb.create_sheet("Not Interested")
    _write_table(ws, hd, "NotInterested", money=("Price",), miles=("Mileage",),
                 links=("Seller Listing URL", "Aggregator URL"), widths={"Note": 40, "Seller Listing URL": 40})

    # 8. Assumptions (so every number above can be traced)
    ws = wb.create_sheet("Assumptions")
    f = settings["finance"]
    assumptions = [
        ("Exported", _now_local()), ("Home ZIP", settings.get("home_zip")),
        ("Rows in Current Inventory", "All vehicles" if settings["export"].get("include_filtered_out") else "Vehicles matching current filters"),
        ("Target monthly payment", f["target_monthly"]), ("Loan terms (months)", ", ".join(map(str, f["terms"]))),
        ("APR scenarios (%)", ", ".join(map(str, f["aprs"]))), ("Current vehicle", f.get("current_vehicle")),
        ("Current loan payoff", f["trade_payoff"]), ("Trade-in estimate (low)", f["trade_in_low"]),
        ("Trade-in estimate (high)", f["trade_in_high"]), ("Separate cash offer", f["cash_offer"]),
        ("Sales tax rate (%)", f["sales_tax_rate"]), ("DMV fees", f["dmv_fees"]), ("Dealer doc fee", f["doc_fee"]),
        ("Extra cash down", f["extra_cash_down"]),
        ("Dealer trade-in reduces taxable price", f["trade_in_reduces_taxable"]),
        ("Separate cash sale reduces taxable price", f["cash_sale_reduces_taxable"]),
        ("Primary scenario (inventory sheet)", f"{finance.SCENARIO_LABELS.get(f['primary_scenario'])}, {f['primary_term']} mo @ {f['primary_apr']}%"),
        ("Payment formula", "P × r / (1 − (1 + r)^−n), r = APR / 12"),
        ("Distances", "Straight-line miles from home ZIP (not driving distance)"),
        ("Warranty", "Rivian terms use the in-service date; other brands need the truck's actual warranty dates and mileage limits"),
        ("Tow capacity", "Source-backed or user-verified capacity for the truck; unknown is not inferred from model maximums"),
        ("Powered tonneau", "'Confirmed installed & working' is only set by your own verification"),
        ("Missing vs removed", "'missing' = not seen in the latest complete scan; 'removed' = confirmed by you or the provider"),
        ("Seller vs aggregator", "'Seller Listing URL' is the page of whoever holds the truck (dealer site, Rivian, CarMax, "
                                 "Carvana, private ad); 'Aggregator URL' is a re-listing (AutoTrader, Cars.com, CarGurus…)"),
        ("Not interested", "Trucks you hid are excluded from Current Inventory and listed on the 'Not Interested' sheet"),
        ("Deal rating", "Price vs. a regression on your own collected listings (year, miles, battery, motor, tonneau); "
                        "branded-title/accident trucks excluded; ±3% = fair, ≤ −8% = great deal, > +8% = overpriced"),
        ("Negotiation leverage", "Heuristic: High = listed 60+ days or 2+ price cuts; Medium = 30+ days or 1 cut"),
    ]
    _write_table(ws, pd.DataFrame(assumptions, columns=["Assumption", "Value"]), "Assumptions",
                 widths={"Assumption": 42, "Value": 80}, style="TableStyleLight9")
    for row in ws.iter_rows(min_row=2):
        if isinstance(row[1].value, float) and row[0].value and "rate" not in row[0].value.lower() and "%" not in row[0].value:
            row[1].number_format = MONEY
    return wb


def _subset(df: pd.DataFrame, mapping: dict) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=list(mapping.values()))
    out = df[[c for c in mapping if c in df.columns]].rename(columns=mapping)
    for col in ("Current Price", "Mileage"):
        if col in out:
            out[col] = pd.to_numeric(out[col], errors="coerce")
    if "Demo Data" in out:
        out["Demo Data"] = out["Demo Data"].fillna(0).astype(bool)
    return out


def export_bytes(conn: sqlite3.Connection, settings: dict) -> bytes:
    buf = io.BytesIO()
    build_workbook(conn, settings).save(buf)
    return buf.getvalue()


def export_file(conn: sqlite3.Connection, settings: dict, directory: Path) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"rivian_finder_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    build_workbook(conn, settings).save(path)
    return path

"""Seller-vs-aggregator links, 'not interested' hiding, and user-added seller links."""
import openpyxl
import pytest
from conftest import FakeAdapter, listing, vin

from rivian_finder import attributes, dedup, export_xlsx, hiding, links, pipeline, queries, services
from rivian_finder.models import FetchResult


@pytest.mark.parametrize("url,seller_type,role", [
    ("https://www.autotrader.com/cars-for-sale/vehicle/711000001", None, "aggregator"),
    ("https://www.cars.com/vehicledetail/abc-123/", "dealer", "aggregator"),
    ("https://www.cargurus.com/Cars/inventorylisting/x#listing=123456789", None, "aggregator"),
    ("https://www.autotempest.com/results?make=rivian", None, "aggregator"),
    ("https://www.autotrader.com/cars-for-sale/vehicle/711000009", "private", "seller"),  # private-party ad
    ("https://www.carvana.com/vehicle/3456789", None, "seller"),
    ("https://www.carmax.com/car/27612345", None, "seller"),
    ("https://rivian.com/configurations/list?INVENTORY_TYPE=PRE_OWNED_VEHICLE", None, "seller"),
    ("https://cars.ksl.com/listing/9900001", "private", "seller"),
    ("https://www.somedealer-ev.com/used/2024-rivian-r1t", "dealer", "seller"),
    (None, None, "none"),
])
def test_listing_role(url, seller_type, role):
    assert links.listing_role(url, seller_type)[0] == role


def test_seller_link_beats_aggregator_for_same_vin(conn, settings):
    pipeline.import_listings(conn, [listing("autotrader", import_method="email", vin=vin(1),
                                            url="https://www.autotrader.com/cars-for-sale/vehicle/711000001",
                                            source_listing_id="711000001")], "email_import")
    pipeline.run_search(conn, [FakeAdapter("autodev", [FetchResult(
        [listing(vin=vin(1), url="https://www.somedealer-ev.com/used/" + vin(1), seller_type="dealer")], True, "s")])], {})
    row = queries.inventory_frame(conn, settings).iloc[0]
    assert row["seller_url"].startswith("https://www.somedealer-ev.com/")
    assert row["seller_site"].startswith("Dealer site")
    assert row["found_on"] == "AutoTrader"
    assert row["link_status"] == "seller"
    assert row["find_seller_url"] == ""


def test_aggregator_only_gets_vin_search_link(conn, settings):
    pipeline.import_listings(conn, [listing("cars_com", import_method="email", vin=vin(2),
                                            url="https://www.cars.com/vehicledetail/abc-123/")], "email_import")
    row = queries.inventory_frame(conn, settings).iloc[0]
    assert row["seller_url"] == ""
    assert row["link_status"] == "aggregator only"
    assert vin(2) in row["find_seller_url"] and row["find_seller_url"].startswith("https://www.google.com/search")


def test_add_seller_link_to_vinless_truck(conn, settings):
    pipeline.import_listings(conn, [listing("cargurus", import_method="email", vin=None, price=58000,
                                            url="https://www.cargurus.com/Cars/x#listing=123456789")], "email_import")
    vid = conn.execute("SELECT id FROM vehicles").fetchone()[0]
    out = services.add_listing_link(conn, vid, "https://www.dealer-two.com/inventory/r1t-55", price=57500)
    assert out["vehicle_id"] == vid and out["created"]
    assert conn.execute("SELECT COUNT(*) FROM vehicles").fetchone()[0] == 1
    row = queries.inventory_frame(conn, settings).iloc[0]
    assert row["seller_url"] == "https://www.dealer-two.com/inventory/r1t-55"
    with pytest.raises(ValueError):
        services.add_listing_link(conn, vid, "not a url")


def test_marketcheck_dealer_website_is_kept(conn, settings):
    pipeline.run_search(conn, [FakeAdapter("marketcheck", [FetchResult(
        [listing("marketcheck", sid="m1", vin=vin(3), url="https://www.autotrader.com/cars-for-sale/vehicle/7",
                 dealer_website="https://dealer-three.example")], True, "s")])], {})
    assert queries.inventory_frame(conn, settings).iloc[0]["dealer_website"] == "https://dealer-three.example"


def _two_trucks(conn):
    a = FakeAdapter("autodev", [FetchResult([listing(vin=vin(1)), listing(vin=vin(2))], True, "s"),
                                FetchResult([listing(vin=vin(1), price=55000), listing(vin=vin(2), price=55000)], True, "s")])
    pipeline.run_search(conn, [a], {})
    return a


def test_hide_and_restore(conn, settings):
    _two_trucks(conn)
    v1 = conn.execute("SELECT id FROM vehicles WHERE vin=?", (vin(1),)).fetchone()[0]
    hiding.hide_vehicle(conn, v1, "other", "too pricey")
    inv = queries.inventory_frame(conn, settings)
    f = settings["filters"]
    assert v1 not in set(queries.apply_filters(inv, f)["vehicle_id"])
    assert set(queries.apply_filters(inv, dict(f, hidden_only=True))["vehicle_id"]) == {v1}
    assert v1 in set(queries.apply_filters(inv, dict(f, show_hidden=True))["vehicle_id"])
    hiding.restore_vehicle(conn, v1)
    assert v1 in set(queries.apply_filters(queries.inventory_frame(conn, settings), f)["vehicle_id"])


def test_hide_reason_records_the_fact(conn):
    _two_trucks(conn)
    v1 = conn.execute("SELECT id FROM vehicles WHERE vin=?", (vin(1),)).fetchone()[0]
    hiding.hide_vehicle(conn, v1, "title_issues", "Carfax: rebuilt")
    r = attributes.resolve(conn, v1)["title_status"]
    assert r.value == "branded" and r.confidence == "verified"
    hiding.restore_vehicle(conn, v1)
    assert attributes.resolve(conn, v1)["title_status"].value == "branded"   # the fact survives restoring


def test_hide_as_sold_marks_listings_removed(conn):
    _two_trucks(conn)
    v1 = conn.execute("SELECT id FROM vehicles WHERE vin=?", (vin(1),)).fetchone()[0]
    hiding.hide_vehicle(conn, v1, "sold")
    assert conn.execute("SELECT status FROM listings WHERE vehicle_id=?", (v1,)).fetchone()[0] == "removed"


def test_hidden_trucks_excluded_from_run_counts(conn):
    a = _two_trucks(conn)
    v1 = conn.execute("SELECT id FROM vehicles WHERE vin=?", (vin(1),)).fetchone()[0]
    hiding.hide_vehicle(conn, v1, "too_far")
    s = pipeline.run_search(conn, [a], {})
    assert s["counts"]["price_drop"] == 1
    assert s["counts"]["hidden_vehicle_updates"] == 1
    assert "marked not interested" in queries.summary_text(s, 1)


def test_hidden_truck_stays_hidden_when_relisted_elsewhere(conn, settings):
    _two_trucks(conn)
    v1 = conn.execute("SELECT id FROM vehicles WHERE vin=?", (vin(1),)).fetchone()[0]
    hiding.hide_vehicle(conn, v1, "other")
    pipeline.import_listings(conn, [listing("ksl", import_method="email", vin=vin(1),
                                            url="https://cars.ksl.com/listing/9900099", source_listing_id="9900099")],
                             "email_import")
    inv = queries.inventory_frame(conn, settings)
    assert v1 not in set(queries.apply_filters(inv, settings["filters"])["vehicle_id"])


def test_merge_keeps_hidden_flag(conn):
    pipeline.import_listings(conn, [listing("ksl", import_method="manual", url="https://x.example/1", vin=None),
                                    listing("dealer", import_method="manual", url="https://x.example/2", vin=vin(8))],
                             "manual_entry")
    a, b = [r[0] for r in conn.execute("SELECT id FROM vehicles ORDER BY id")]
    hiding.hide_vehicle(conn, a, "other")
    dedup.merge_vehicles(conn, b, a)
    assert conn.execute("SELECT hidden_reason FROM vehicles WHERE id=?", (b,)).fetchone()[0] == "other"


def test_export_has_seller_links_and_not_interested_sheet(conn, settings, tmp_path):
    _two_trucks(conn)
    v1 = conn.execute("SELECT id FROM vehicles WHERE vin=?", (vin(1),)).fetchone()[0]
    hiding.hide_vehicle(conn, v1, "title_issues")
    wb = openpyxl.load_workbook(export_xlsx.export_file(conn, settings, tmp_path))
    inv = wb["Current Inventory"]
    headers = [c.value for c in inv[1]]
    assert {"Sold By", "Seller Listing URL", "Found On (aggregators)", "Aggregator URL"} <= set(headers)
    vins = [inv.cell(row=r, column=headers.index("VIN") + 1).value for r in range(2, inv.max_row + 1)]
    assert vin(1) not in vins and vin(2) in vins
    ni = wb["Not Interested"]
    assert ni.cell(row=2, column=[c.value for c in ni[1]].index("Reason") + 1).value.startswith("Title issues")


def test_display_name_is_just_who_it_is():
    assert links.display_name("anchor auto outlet inc") == "Anchor Auto Outlet Inc"
    assert links.display_name("CARMAX SALT LAKE") == "CarMax Salt Lake"
    assert links.display_name("AutoMatch") == "AutoMatch"
    assert links.display_name("", "Dealer site (automatchutah.com)") == "automatchutah.com"
    assert links.display_name(None, "Rivian (manufacturer)") == "Rivian"


def test_inventory_sold_by_prefers_the_seller_listing(conn, settings):
    pipeline.import_listings(conn, [listing("autotrader", import_method="email", vin=vin(1), seller_name=None,
                                            url="https://www.autotrader.com/cars-for-sale/vehicle/711000001",
                                            source_listing_id="711000001")], "email_import")
    pipeline.run_search(conn, [FakeAdapter("autodev", [FetchResult(
        [listing(vin=vin(1), url="https://www.somedealer-ev.com/used/x", seller_name="some dealer ev",
                 seller_type="dealer")], True, "s")])], {})
    assert queries.inventory_frame(conn, settings).iloc[0]["sold_by"] == "Some Dealer EV"

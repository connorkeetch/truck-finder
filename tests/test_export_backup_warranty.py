import json
from datetime import date

import openpyxl
from conftest import FakeAdapter, listing, vin

from rivian_finder import backup, checklist, db, export_xlsx, pipeline, queries, services, warranty
from rivian_finder.models import FetchResult
from rivian_finder.sources.demo import DemoAdapter

TERMS = {"comprehensive_years": 5, "comprehensive_miles": 60000, "battery_years": 8,
         "battery_miles_large_max": 175000, "battery_miles_standard": 150000}


def test_warranty_unknown_without_in_service_date():
    w = warranty.remaining(None, 20000, TERMS)
    assert not w["known"] and "Unknown" in w["summary"]


def test_warranty_time_and_miles():
    w = warranty.remaining("2023-06-15", 20000, TERMS, "Large", today=date(2026, 10, 20), motor="Quad")
    assert w["comprehensive_expires"] == "2028-06-15"
    assert w["comprehensive_months_left"] == 19
    assert w["comprehensive_miles_left"] == 40000
    assert w["battery_expires"] == "2031-06-15" and w["battery_miles_left"] == 155000


def test_warranty_expires_by_miles_first():
    w = warranty.remaining("2025-01-01", 61000, TERMS, today=date(2026, 1, 1))
    assert w["comprehensive_months_left"] == 0 and w["summary"] == "Comprehensive expired"


def test_standard_pack_battery_miles():
    assert warranty.remaining("2025-01-01", 0, TERMS, "Standard", today=date(2025, 1, 2))["battery_miles_left"] == 150000


def _populate(conn):
    a = DemoAdapter(conn)
    for _ in range(3):
        pipeline.run_search(conn, [a], {})


def test_inventory_and_filters(conn, settings):
    _populate(conn)
    inv = queries.inventory_frame(conn, settings)
    assert not inv.empty
    f = dict(settings["filters"], near_misses=False)
    out = queries.apply_filters(inv, f)
    assert (out["price"].dropna() <= f["max_price"]).all()
    assert (out["miles"].dropna() <= f["max_miles"]).all()
    assert set(out["battery"]) <= set(f["batteries"])
    assert (out["status"] == "active").all()
    f2 = dict(f, show_demo=False)
    assert queries.apply_filters(inv, f2).empty


def test_summary_text_format():
    s = {"counts": {"new_listing": 8, "price_drop": 3, "disappeared": 4}, "sources": [
        {"source": "carvana", "name": "Carvana", "status": "failed", "error": "HTTP 503"}]}
    text = queries.summary_text(s, 127)
    assert text.startswith("127 matching vehicles, 8 new, 3 price reductions, 4 missing since last successful scan.")
    assert "Carvana failed (HTTP 503) — previous listings left unchanged." in text


def test_excel_export_structure(conn, settings, tmp_path):
    _populate(conn)
    path = export_xlsx.export_file(conn, settings, tmp_path)
    wb = openpyxl.load_workbook(path)
    expected = ["Current Inventory", "New Listings", "Price Changes", "Missing - Removed", "Listing History",
                "Source Status", "Financial Comparisons", "Not Interested", "Assumptions"]
    assert wb.sheetnames == expected
    for name in expected:
        ws = wb[name]
        assert len(ws.tables) == 1, name
        assert ws.freeze_panes in ("A2", "B2"), name
    inv = wb["Current Inventory"]
    headers = [c.value for c in inv[1]]
    for col in ("VIN", "Listing URL", "Captured / Last Seen (UTC)", "Seller", "Source", "Your Price (after discounts)",
                "Warranty Ends", "EPA Range (mi)"):
        assert col in headers
    for gone in ("Powered Tonneau", "Tonneau Source", "Demo Data", "Wheels", "Title", "Accidents", "Checklist",
                 "Last Contact (UTC)", "Follow-up", "Owners", "Recalls (model year)"):
        assert gone not in headers, gone
    price_col = headers.index("Price") + 1
    assert inv.cell(row=2, column=price_col).number_format == '"$"#,##0'
    assert inv.conditional_formatting
    fin_headers = [c.value for c in wb["Financial Comparisons"][1]]
    assert "Payment 72 mo @ 4.99%" in fin_headers and "Cash for $800/mo 60 mo @ 4.49%" in fin_headers


def test_export_with_empty_database(conn, settings, tmp_path):
    assert export_xlsx.export_file(conn, settings, tmp_path).exists()


def test_backup_restore_and_json(tmp_path):
    live = tmp_path / "live.db"
    c = db.connect(live)
    pipeline.run_search(c, [FakeAdapter("autodev", [FetchResult([listing(vin=vin(1))], True, "s")])], {})
    bpath = backup.backup_db(c, tmp_path / "b")
    payload = backup.export_json(c)
    c.close()
    fresh = db.connect(tmp_path / "fresh.db")
    counts = backup.import_json(fresh, payload)
    assert counts["vehicles"] == 1 and counts["observations"] == 1
    assert fresh.execute("SELECT vin FROM vehicles").fetchone()[0] == vin(1)
    fresh.close()
    backup.restore_db(bpath, live)
    c = db.connect(live)
    assert c.execute("SELECT COUNT(*) FROM listings").fetchone()[0] == 1


def test_json_export_covers_schema_and_preserves_shopping_history(conn):
    with conn:
        vehicle_id = conn.execute("INSERT INTO vehicles(vin, created_at, updated_at) VALUES (?,?,?)",
                                  (vin(1), db.utcnow(), db.utcnow())).lastrowid
    checklist.save(conn, vehicle_id, {"tires": ("problem", "Rear tires need replacing")})
    checklist.add_contact(conn, vehicle_id, "email", "Example seller", "Requested an itemized quote",
                          follow_up=date(2026, 11, 1))
    payload = backup.export_json(conn)
    exported = json.loads(payload)["tables"]
    tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")}
    assert set(exported) == tables - {"schema_migrations"}
    with db.connect(":memory:") as fresh:
        counts = backup.import_json(fresh, payload)
        assert counts["checklist"] == counts["contacts"] == 1
        assert checklist.get(fresh, vehicle_id)["tires"] == checklist.get(conn, vehicle_id)["tires"]
        assert checklist.contacts(fresh, vehicle_id) == checklist.contacts(conn, vehicle_id)
        assert not fresh.execute("PRAGMA foreign_key_check").fetchall()


def test_json_import_accepts_older_exports_without_newer_tables(conn):
    exported = json.loads(backup.export_json(conn))
    for table in ("api_usage", "recalls", "checklist", "contacts", "link_checks", "email_inbox"):
        del exported["tables"][table]
    with db.connect(":memory:") as fresh:
        counts = backup.import_json(fresh, json.dumps(exported))
        assert counts["contacts"] == counts["checklist"] == counts["email_inbox"] == 0
        assert fresh.execute("SELECT COUNT(*) FROM sources").fetchone()[0] == len(exported["tables"]["sources"])


def test_purge_demo(conn):
    _populate(conn)
    n = services.purge_demo(conn)
    assert n > 20
    assert conn.execute("SELECT COUNT(*) FROM vehicles").fetchone()[0] == 0
    assert conn.execute("SELECT COUNT(*) FROM listings").fetchone()[0] == 0


def test_migrations_are_idempotent(tmp_path):
    p = tmp_path / "m.db"
    db.connect(p).close()
    c = db.connect(p)
    assert db.migrate(c) == []
    assert c.execute("SELECT COUNT(*) FROM saved_searches").fetchone()[0] > 5
